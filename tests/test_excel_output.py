"""Regional workbook layout (src/excel_output.py), on synthetic summaries (no network)."""

import pandas as pd
from openpyxl import load_workbook

from src.excel_output import generate_all_workbooks

REGION_ROWS = [
    # DUID, FY25-26, FY26-27, FY26-27 Import, draft, YoY
    ("BESS1", 0.97, 0.96, 1.02, 0.90, -0.01),
    ("BESS2", 0.95, 0.98, 0.99, 0.99, 0.03),
    ("SOLAR1", 0.90, 0.88, None, 0.97, -0.02),
    ("WIND1", 0.99, 0.99, None, 0.95, 0.0),
    ("WIND2", 0.96, 0.96, None, 0.96, 0.0),
    ("HYDRO1", 1.00, 1.01, None, 1.00, 0.01),
    ("BESS3", 0.94, 0.95, 1.03, 0.92, 0.01),
    ("BESS4", 0.93, 0.93, 0.97, 0.93, 0.0),
    ("BESS5", 0.92, 0.91, 1.04, 0.94, -0.01),   # five import MLFs: the workbook's threshold for an Import column
]


def summary(region="TAS1") -> pd.DataFrame:
    rows = []
    for duid, prev, cur, imp, draft, yoy in REGION_ROWS:
        rows.append({"DUID": duid, "REGIONID": region, "STATION_NAME": duid, "FUEL_CATEGORY": "", "CAPACITY_MW": 10,
                     "FY25-26": prev, "FY26-27": cur, "FY26-27 Import": imp, "FY27-28 (Draft)": draft,
                     "LATEST_MLF": cur, "PREV_MLF": prev, "YOY_CHANGE": yoy, "YOY_PCT_CHANGE": yoy / prev * 100})
    return pd.DataFrame(rows)


def workbook(tmp_path):
    generate_all_workbooks(summary(), str(tmp_path))
    return load_workbook(tmp_path / "TAS_mlf.xlsx")


def header(ws) -> list:
    return [c.value for c in ws[1]]


def test_heatmap_scale_skips_import_columns(tmp_path):
    ws = workbook(tmp_path)["Heatmap"]
    cols = header(ws)
    scaled = {str(cf.sqref).split(":")[0].rstrip("0123456789") for cf in ws.conditional_formatting}
    letters = {name: ws.cell(row=1, column=i + 1).column_letter for i, name in enumerate(cols)}
    assert letters["FY26-27"] in scaled and letters["FY25-26"] in scaled
    assert letters["FY26-27 Import"] not in scaled


def movers(ws) -> dict[str, list[str]]:
    out, section = {}, None
    for row in ws.iter_rows(values_only=True):
        if isinstance(row[0], str) and "— Most " in row[0]:
            section = "degraded" if "Degraded" in row[0] else "improved"
            out[section] = []
        elif section and row[0] not in (None, "DUID"):
            out[section].append(row[0])
    return out


def test_movers_list_only_units_that_moved_that_way(tmp_path):
    m = movers(workbook(tmp_path)["Biggest Movers"])
    assert m["degraded"] == ["SOLAR1", "BESS1", "BESS5"]          # not WIND1/WIND2/BESS4 (0.0000) or improvers
    assert m["improved"] == ["BESS2", "HYDRO1", "BESS3"]


def test_tables_sort_by_latest_final_year_not_the_draft(tmp_path):
    wb = workbook(tmp_path)
    for sheet in ("MLF Table", "Heatmap"):
        first = [wb[sheet].cell(row=r, column=1).value for r in (2, 3)]
        assert first == ["SOLAR1", "BESS5"], sheet                # FY26-27 0.88, 0.91 — not draft BESS1 0.90
